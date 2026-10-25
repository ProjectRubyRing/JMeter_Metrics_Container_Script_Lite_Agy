# -*- coding: utf-8 -*-
"""
create_metrics_reference_excel.py
JMeterメトリクス収集・コンテナ性能監視スクリプトにおける
「レポート出力項目一覧」および「そのまま実装に利用できるコマンドリファレンス」を
Meiryo UIフォント・モノトーン配色でExcelブックとして生成するスクリプト。
"""

import openpyxl
import openpyxl.styles as styles
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# デフォルトフォントを Meiryo UI に変更 (全セルの完全統一)
styles.DEFAULT_FONT.name = "Meiryo UI"

def build_excel(output_path):
    wb = openpyxl.Workbook()
    default_sheet = wb.active

    # Normal NamedStyle のフォントを Meiryo UI に設定
    for ns in wb._named_styles:
        if ns.name == "Normal":
            ns.font = styles.DEFAULT_FONT

    # -------------------------------------------------------------
    # スタイル定義 (Meiryo UI & モノトーンベース)
    # -------------------------------------------------------------
    FONT_NAME = "Meiryo UI"

    f_title = Font(name=FONT_NAME, size=13, bold=True, color="FFFFFF")
    f_sub = Font(name=FONT_NAME, size=9, italic=True, color="E0E0E0")
    f_section = Font(name=FONT_NAME, size=10.5, bold=True, color="000000")
    f_th = Font(name=FONT_NAME, size=9.5, bold=True, color="FFFFFF")
    f_data = Font(name=FONT_NAME, size=9, color="1A1A1A")
    f_data_bold = Font(name=FONT_NAME, size=9, bold=True, color="000000")
    f_cmd = Font(name=FONT_NAME, size=8.5, color="111111")
    f_badge = Font(name=FONT_NAME, size=9, bold=True, color="FFFFFF")
    f_small = Font(name=FONT_NAME, size=8.5, color="444444")

    # 背景色（モノトーン階調）
    fill_title = PatternFill(start_color="242424", end_color="242424", fill_type="solid")       # ほぼ黒
    fill_th = PatternFill(start_color="444444", end_color="444444", fill_type="solid")          # ダークチャコール
    fill_row_even = PatternFill(start_color="F7F7F7", end_color="F7F7F7", fill_type="solid")    # 極薄グレー
    fill_row_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")     # 白
    fill_cmd = PatternFill(start_color="EFEFEF", end_color="EFEFEF", fill_type="solid")         # コード背景薄グレー
    fill_cmd_alt = PatternFill(start_color="E6E6E6", end_color="E6E6E6", fill_type="solid")     # コード背景やや濃いグレー

    # バッジ用
    fill_b_k1 = PatternFill(start_color="2E2E2E", end_color="2E2E2E", fill_type="solid")        # ① jboss-cli
    fill_b_k2 = PatternFill(start_color="454545", end_color="454545", fill_type="solid")        # ② jstat/jcmd/JFR
    fill_b_k3 = PatternFill(start_color="5C5C5C", end_color="5C5C5C", fill_type="solid")        # ③ top/ps/proc
    fill_b_k4 = PatternFill(start_color="1A1A1A", end_color="1A1A1A", fill_type="solid")        # ④ CPU/cgroup
    fill_b_cmn = PatternFill(start_color="4D4D4D", end_color="4D4D4D", fill_type="solid")       # 共通/レポート

    # 罫線
    c_border_light = "D8D8D8"
    c_border_mid = "8C8C8C"
    c_border_dark = "333333"

    bd_thin = Border(
        left=Side(style='thin', color=c_border_light),
        right=Side(style='thin', color=c_border_light),
        top=Side(style='thin', color=c_border_light),
        bottom=Side(style='thin', color=c_border_light)
    )
    bd_th = Border(
        left=Side(style='thin', color=c_border_mid),
        right=Side(style='thin', color=c_border_mid),
        top=Side(style='medium', color=c_border_dark),
        bottom=Side(style='medium', color=c_border_dark)
    )

    # アラインメント
    al_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    al_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    al_right = Alignment(horizontal="right", vertical="center", wrap_text=True)
    al_cmd = Alignment(horizontal="left", vertical="center", wrap_text=True)

    fill_badges = {"①": fill_b_k1, "②": fill_b_k2, "③": fill_b_k3, "④": fill_b_k4, "共通": fill_b_cmn}

    # =============================================================
    # 補助関数
    # =============================================================
    def apply_title_banner(ws, title_text, sub_text, max_col):
        ws.row_dimensions[1].height = 26
        ws.row_dimensions[2].height = 18
        
        for r in range(1, 3):
            for c in range(1, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = fill_title
                cell.font = Font(name=FONT_NAME, size=9)
        
        c1 = ws.cell(row=1, column=1, value=title_text)
        c1.font = f_title
        c1.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        c2 = ws.cell(row=2, column=1, value=sub_text)
        c2.font = f_sub
        c2.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)

    def merge_and_style(ws, start_row, start_col, end_row, end_col, font=None, fill=None, alignment=None, border=None):
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                cell = ws.cell(row=r, column=c)
                if font: cell.font = font
                else: cell.font = Font(name=FONT_NAME, size=9)
                if fill: cell.fill = fill
                if alignment: cell.alignment = alignment
                if border: cell.border = border
        if start_row != end_row or start_col != end_col:
            ws.merge_cells(start_row=start_row, start_column=start_col, end_row=end_row, end_column=end_col)

    def auto_fit_columns(ws, min_widths=None, max_widths=None):
        if min_widths is None: min_widths = {}
        if max_widths is None: max_widths = {}
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                if cell.row in [1, 2]: continue
                v = str(cell.value or '')
                lines = v.split('\n')
                for line in lines:
                    length = 0
                    for ch in line:
                        length += 2 if ord(ch) > 127 else 1
                    if length > max_len:
                        max_len = length
            calc_w = max(max_len * 1.05 + 3, 10)
            if col_letter in min_widths:
                calc_w = max(calc_w, min_widths[col_letter])
            if col_letter in max_widths:
                calc_w = min(calc_w, max_widths[col_letter])
            ws.column_dimensions[col_letter].width = round(calc_w, 1)

    def enforce_meiryo_ui(workbook):
        for sheet in workbook.worksheets:
            for r in range(1, sheet.max_row + 1):
                for c in range(1, sheet.max_column + 1):
                    cell = sheet.cell(row=r, column=c)
                    if cell.font is None:
                        cell.font = Font(name=FONT_NAME, size=9)
                    elif cell.font.name != FONT_NAME:
                        f = cell.font
                        cell.font = Font(
                            name=FONT_NAME,
                            size=f.size,
                            bold=f.bold,
                            italic=f.italic,
                            color=f.color,
                            vertAlign=f.vertAlign,
                            underline=f.underline,
                            strike=f.strike
                        )

    # =============================================================
    # SHEET 1: 目次・全体構成 (Overview)
    # =============================================================
    ws1 = wb.create_sheet(title="目次・全体構成")
    apply_title_banner(ws1, "JMeter 性能試験コンテナメトリクス収集 — 全体構成・設計概要", 
                       "ECS Fargate (UBI 9.8 + OpenJDK 21 + JBoss EAP 8.1) 内部可観測性スクリプト & Excel レポート仕様", 7)
    ws1.views.sheetView[0].showGridLines = True
    ws1.row_dimensions[3].height = 10

    # 1. 概要
    ws1.cell(row=4, column=1, value="1. 収集の目的と CloudWatch Container Insights (拡張) との役割分担").font = f_section
    merge_and_style(ws1, 4, 1, 4, 7)

    intro_data = [
        ("目的", "JMeter 等の負荷試験において、「コンテナの外側」からでは見えない JBoss EAP 8.1 / JVM / Linux カーネル内部のボトルネック（DB接続待ち・GC停止・スレッド枯渇・CPU取り合い）を低負荷 (CPU 1.7%以下) に定周期採取し、Windows PowerShell でモノトーンの Excel レポート (最大29シート・49グラフ) に自動集計する。"),
        ("対象環境", "AWS ECS on Fargate (microVM), OS: Red Hat UBI 9.8 (Minimal), Runtime: Red Hat build of OpenJDK 21 (LTS), AP Server: JBoss EAP 8.1, ストレージ: EFS (readonlyRootFilesystem=true 対応)"),
        ("外側で見えるもの\n(Container Insightsで十分)", "タスク / コンテナの CPU 使用率・使用量、メモリ使用量 (RSS)、ネットワーク送受信バイト数、ストレージ転送量、タスク再起動回数。\n⇒ ※これらは perf-metrics.sh では意図的に重複収集せず、Container Insights の標準グラフに委ねる。"),
        ("内側で見えるもの\n(本スクリプトで収集)", "① DB接続プールの使用率・待ちスレッド数・取得失敗、Undertow リクエスト処理時間、セッション数、IOワーカー/EEスレッドプール負荷\n② JVMヒープ各領域 (Eden/Survivor/Old)、GC停止時間 (最長/合計)、GC原因、JIT時間、コードキャッシュ満杯、NMTネイティブメモリ分類、用途別スレッド状態 (DB待機等)\n③ プロセス別CPU内訳 (JBoss vs 収集負荷)、JBossスレッドの用途別CPU (HTTP/GC/JIT)、CPU上位10スレッド、FD使用率、TCP接続状態 (CLOSE_WAIT/TIME_WAIT)\n④ Fargate microVM とコンテナのCPU取り分不一致による -XX:ActiveProcessorCount の判定材料、CPU steal、PSI (待ち時間割合)、CFSスロットリング")
    ]

    r_idx = 5
    for item, desc in intro_data:
        ws1.row_dimensions[r_idx].height = 42 if '\n' in desc else 24
        ws1.cell(row=r_idx, column=1, value=item)
        ws1.cell(row=r_idx, column=2, value=desc)
        merge_and_style(ws1, r_idx, 1, r_idx, 1, font=f_data_bold, fill=fill_row_even, alignment=al_center, border=bd_thin)
        merge_and_style(ws1, r_idx, 2, r_idx, 7, font=f_data, alignment=al_left, border=bd_thin)
        r_idx += 1

    r_idx += 1
    # 2. 4区分
    ws1.cell(row=r_idx, column=1, value="2. メトリクス収集の 4 区分と採取アーキテクチャ").font = f_section
    merge_and_style(ws1, r_idx, 1, r_idx, 7)
    r_idx += 1

    headers_k = ["区分", "主要ツール", "採取対象・何が分かるか", "間隔", "出力CSVファイル", "出力先シート", "設計ポイント・オーバーヘッド対策"]
    ws1.row_dimensions[r_idx].height = 24
    for i, h in enumerate(headers_k, start=1):
        cell = ws1.cell(row=r_idx, column=i, value=h)
        cell.font = f_th; cell.fill = fill_th; cell.alignment = al_center; cell.border = bd_th
    r_idx += 1

    cat_rows = [
        ("①", "jboss-cli", "DB接続プール統計・情報 (IronJacamar)、HTTP リスナー (Undertow)、HTTPセッション、スレッドプール (EJB/EE)、IOワーカー (タスク/IO)、トランザクション (Narayana)、JVM状態 (MBean)",
         "10秒", "cli_datasource.csv\ncli_datasource_info.csv\ncli_undertow.csv\ncli_session.csv\ncli_threadpool.csv\ncli_io_worker.csv\ncli_transactions.csv\ncli_jvm.csv",
         "①DB接続プール\n①HTTP\n①セッション\n①スレッドプール\n①IOワーカー\n①トランザクション\n①JVM（CLI）",
         "常駐プロセスを1つだけ維持し、FIFOパイプでDMRクエリを一括バッチ送信。応答DMRをawkで高速パースし、CLI再起動オーバーヘッドを徹底排除。"),
        ("②", "jstat\njcmd\nJFR", "ヒープ各領域 (Eden/Survivor/Old)、GC回数・時間・直近原因、セーフポイント停止、JIT時間、G1リージョン、コードキャッシュ、NMT (Native Memory Tracking)、用途別スレッド状態、JFRイベント (GC停止・Hot Method・Lock競合・IO)",
         "jstat: 10秒\njcmd: 60秒\nJFR: 常時記録", "jvm_jstat_gc.csv\njvm_jcmd_heap.csv\njvm_codecache.csv\njvm_nmt.csv\njvm_thread_states.csv\njfr_gc.csv, jfr_stats.csv\njfr_hot_methods.csv\njfr_contention.csv\njfr_io.csv, jfr_summary.csv",
         "②ヒープ・GC（jstat）\n②ヒープ（jcmd）\n②コードキャッシュ\n②NMT\n②スレッド状態\n②GC（JFR）\n②JVM統計（JFR）\n②JFR上位\n②JFR記録",
         "jstatはカスタム perfgc 定義で1プロセス常駐。jcmdはJVM起動が伴うため60秒間隔。JFRは起動時オプションでヒープ・GC・統計に限定し低負荷化。試験後にJava単一ソース (PerfJfrCsv.java) でCSV化。"),
        ("③", "top\nps\n/proc", "プロセス別CPU (JBoss vs 収集プロセス)、JBossスレッドの用途別CPU (HTTP処理/GC/JIT)、CPU上位10スレッド (TID/nid)、ファイル記述子数 (FD)、TCPソケット状態 (ESTABLISHED/CLOSE_WAIT/TIME_WAIT)、非自発的コンテキストスイッチ",
         "10秒", "os_top_summary.csv\nos_top_process.csv\nos_top_threads.csv\nos_ps.csv\nos_jboss_threads.csv\nos_jboss_proc.csv",
         "③top概要\n③topプロセス\n③topスレッド\n③ps\n③JBossスレッド\n③FD・TCP",
         "UBI最小イメージに procps-ng (top/ps) が無い場合でも /proc/<pid>/task/*/stat, /proc/net/tcp 等から自動フォールバック算出。収集処理自身のCPUを明示し試験の公平性を担保。"),
        ("④", "ECS Metadata\ncgroup\njcmd / CLI", "Fargate microVM 内でのコンテナ取り分 CPU の算出、JVMが認識した CPU 数との比較、-XX:ActiveProcessorCount の要否自動判定、実測裏付け (JBoss実測CPU vs 取り分、GC/JITコア数、実行待ちスレッド、steal、PSI、スロットリング)",
         "判定: 1回\n推移: 10秒", "cpu_env.csv\ncpu_sched.csv",
         "④CPU判定\n④CPUスケジューリング\nサマリー",
         "ECSタスクメタデータv4とJVMのRuntime.availableProcessors()を突合。FargateがCPUシェア配分であることに起因するJVMの過剰スレッド生成・CPU奪い合い問題を自動検出し推奨設定を提示。"),
        ("共通", "エントリポイント\nPowerShell", "収集環境メタデータ (ホスト/OS/JVM/タスク定義)、収集ログ (INFO/WARN/ERROR/RESULT)、Excelレポート生成 (Office Open XML 直接出力、Excel非依存)",
         "起動時/随時", "perf_meta.csv\ncollector.log",
         "サマリー\nグラフ (49グラフ)\n収集ログ\n収集情報\n凡例",
         "readonlyRootFilesystem=true でもEFS上のタスク専用ディレクトリ (pin方式) へ出力。PowerShell 5.1/7対応。Meiryo UIモノトーンで最大29シート・49グラフを5〜50秒で高速生成。")
    ]

    for b, tool, desc, iv, csvs, sheets, note in cat_rows:
        ws1.row_dimensions[r_idx].height = 64
        c_b = ws1.cell(row=r_idx, column=1, value=b)
        c_b.font = f_badge; c_b.fill = fill_badges.get(b, fill_th); c_b.alignment = al_center; c_b.border = bd_thin
        
        c_t = ws1.cell(row=r_idx, column=2, value=tool)
        c_t.font = f_data_bold; c_t.alignment = al_center; c_t.border = bd_thin
        
        c_desc = ws1.cell(row=r_idx, column=3, value=desc)
        c_desc.font = f_data; c_desc.alignment = al_left; c_desc.border = bd_thin
        
        c_iv = ws1.cell(row=r_idx, column=4, value=iv)
        c_iv.font = f_data; c_iv.alignment = al_center; c_iv.border = bd_thin
        
        c_csv = ws1.cell(row=r_idx, column=5, value=csvs)
        c_csv.font = f_small; c_csv.alignment = al_left; c_csv.border = bd_thin
        
        c_sh = ws1.cell(row=r_idx, column=6, value=sheets)
        c_sh.font = f_small; c_sh.alignment = al_left; c_sh.border = bd_thin
        
        c_no = ws1.cell(row=r_idx, column=7, value=note)
        c_no.font = f_data; c_no.alignment = al_left; c_no.border = bd_thin
        r_idx += 1

    r_idx += 1
    # 3. 本ブックの構成
    ws1.cell(row=r_idx, column=1, value="3. 本ドキュメント（本Excelブック）のシート構成").font = f_section
    merge_and_style(ws1, r_idx, 1, r_idx, 7)
    r_idx += 1

    sheet_guides = [
        ("目次・全体構成 (本シート)", "システム概要、4区分の概要、Container Insightsとの対比、本ブックの利用ガイド"),
        ("①レポート出力項目一覧", "New-PerfMetricsReport.ps1 が生成する全29シート・全列・全メトリクスの詳細仕様（項目名、キー、単位、算出式、意味、判定基準）"),
        ("②実装コマンドリファレンス", "各メトリクスを取得する際に「そのまま実装・実行・検証に利用できる」CLI/Linux/Java/AWSコマンド集（ワンライナー・解説付）"),
        ("③サマリー所見・49グラフ定義", "サマリーシートで自動出力される所見（■要対応/▲注意/●参考）の全判定ルール、およびグラフシートの全49グラフ構成"),
        ("④環境設定・タスク定義", "perf-metrics.sh の全環境変数設定、ECS Fargate タスク定義 (JSON) の推奨設定、Docker ビルド引数")
    ]
    for s_name, s_exp in sheet_guides:
        ws1.row_dimensions[r_idx].height = 22
        ws1.cell(row=r_idx, column=1, value=s_name)
        ws1.cell(row=r_idx, column=2, value=s_exp)
        merge_and_style(ws1, r_idx, 1, r_idx, 1, font=f_data_bold, fill=fill_row_even, alignment=al_center, border=bd_thin)
        merge_and_style(ws1, r_idx, 2, r_idx, 7, font=f_data, alignment=al_left, border=bd_thin)
        r_idx += 1

    auto_fit_columns(ws1, min_widths={'A': 16, 'B': 18, 'C': 34, 'D': 14, 'E': 26, 'F': 24, 'G': 38})

    # =============================================================
    # SHEET 2: ①レポート出力項目一覧 (Report Metrics & Columns)
    # =============================================================
    ws2 = wb.create_sheet(title="①レポート出力項目一覧")
    apply_title_banner(ws2, "Excel レポート出力項目・メトリクス定義一覧", 
                       "New-PerfMetricsReport.ps1 が出力する各シートの列定義・算出方法・判定基準・意味の完全リファレンス", 11)
    ws2.views.sheetView[0].showGridLines = True
    ws2.freeze_panes = "E4"
    ws2.row_dimensions[3].height = 26

    cols_sh2 = [
        ("区分", 8, al_center),
        ("シート名", 18, al_center),
        ("ソースCSV", 18, al_left),
        ("カテゴリ", 16, al_left),
        ("項目名 (日本語表示)", 22, al_left),
        ("キー名 / 変数", 20, al_left),
        ("単位", 10, al_center),
        ("種別", 12, al_center),
        ("算出式 / 集計方法", 26, al_left),
        ("項目の意味・目的・解説", 42, al_left),
        ("見方・判定閾値 (要対応■ / 注意▲)", 30, al_left)
    ]

    for ci, (cn, cw, cal) in enumerate(cols_sh2, start=1):
        cell = ws2.cell(row=3, column=ci, value=cn)
        cell.font = f_th; cell.fill = fill_th; cell.alignment = al_center; cell.border = bd_th
        ws2.column_dimensions[get_column_letter(ci)].width = cw

    # 詳細項目データリスト
    metrics_data = [
        # --- ① DB 接続プール ---
        ("①", "①DB接続プール", "cli_datasource.csv", "基本", "データソース", "ds_name", "―", "識別子", "データソース名 (ds_name)", "JBoss に登録されているデータソースの一意な名前", "―"),
        ("①", "①DB接続プール", "cli_datasource.csv", "基本", "種類", "ds_type", "―", "識別子", "ds (標準) / xa (分散XA)", "非XAデータソースかXAデータソースかの識別", "―"),
        ("①", "①DB接続プール", "cli_datasource.csv", "基本", "JNDI 名", "jndi_name", "―", "設定", "jndi-name", "アプリケーションから参照される JNDI 検索名", "―"),
        ("①", "①DB接続プール", "cli_datasource.csv", "接続数", "使用中", "InUseCount", "本", "生値", "InUseCount", "現在アプリケーションスレッドによって貸し出され使用されている物理接続数", "最大プールに達していないか確認"),
        ("①", "①DB接続プール", "cli_datasource.csv", "接続数", "アクティブ", "ActiveCount", "本", "生値", "ActiveCount", "現在 IronJacamar プール内に作成されて存在する物理接続の総数 (使用中 + アイドル)", "最小〜最大プールの範囲で推移"),
        ("①", "①DB接続プール", "cli_datasource.csv", "接続数", "アイドル", "IdleCount", "本", "生値", "IdleCount", "現在空き状態でプールにプールされており、即時貸出可能な接続数", "0 に張り付いている場合は待機発生の予兆"),
        ("①", "①DB接続プール", "cli_datasource.csv", "接続数", "取得可能", "AvailableCount", "本", "生値", "AvailableCount", "最大プール上限まであと何本新規接続を作成可能か", "0 のとき新規接続作成不可"),
        ("①", "①DB接続プール", "cli_datasource.csv", "接続数", "最大プール", "max_pool_size", "本", "設定", "max_pool_size (未設定は20)", "プールが保持・作成できる物理接続の最大許容数", "負荷試験の想定多重度に対して十分か検証"),
        ("①", "①DB接続プール", "cli_datasource.csv", "接続数", "使用率", "InUse÷最大プール", "%", "比率", "InUseCount ÷ max_pool_size", "最大プールに対する使用中接続の割合。データバーで可視化", "▲ 80% 以上で注意、100% で要対応"),
        ("①", "①DB接続プール", "cli_datasource.csv", "待ち・失敗(区間)", "待ち", "Δ WaitCount", "回", "区間増分", "WaitCount - 前回WaitCount", "区間内に接続の空きがなく、スレッドが取得待ち行列に入った回数", "▲ 1回以上で空き待ち発生。接続数不足の疑い"),
        ("①", "①DB接続プール", "cli_datasource.csv", "待ち・失敗(区間)", "取得失敗", "Δ BlockingFailureCount", "回", "区間増分", "BlockingFailureCount - 前回値", "空き接続待ちがタイムアウト (既定30秒) して例外 (IJ000453) になった回数", "■ 1回以上で要対応。アプリで500エラー発生"),
        ("①", "①DB接続プール", "cli_datasource.csv", "待ち・失敗(区間)", "平均待ちスレッド数", "Δ TotalBlockingTime÷経過", "スレッド", "時間比率", "Δ TotalBlockingTime ÷ 経過秒 ÷ 1000", "区間内に「平均して何本のスレッドが電話の前で並んでいたか」を示す実効待機負荷", "▲ 1.0 スレッド以上で恒常的な待ちが発生"),
        ("①", "①DB接続プール", "cli_datasource.csv", "待ち・失敗(区間)", "平均使用接続数", "Δ TotalUsageTime÷経過", "接続", "時間比率", "Δ TotalUsageTime ÷ 経過秒 ÷ 1000", "区間内に実効的に何本の接続が連続使用されていたかの平均値", "トラフィック増加に伴う使用本数の実測傾向確認"),
        ("①", "①DB接続プール", "cli_datasource.csv", "時間(ms,累計)", "平均取得", "AverageGetTime", "ms", "累積平均", "AverageGetTime", "アプリが接続取得を要求してからプールから返却されるまでの平均所要時間", "プール空き時は数ms、枯渇時は急増"),
        ("①", "①DB接続プール", "cli_datasource.csv", "時間(ms,累計)", "最大待ち", "MaxWaitTime", "ms", "累積最大", "MaxWaitTime", "接続取得待ち行列で待たされた最長待機時間 (ms)", "blocking-timeout-wait-millis (30,000ms) 近傍なら危険"),
        ("①", "①DB接続プール", "cli_datasource.csv", "物理接続(区間)", "作成", "Δ CreatedCount", "本", "区間増分", "CreatedCount - 前回値", "区間内にDBサーバーと新規にハンドシェイク・確立した物理TCP/JDBC接続数", "試験開始直後以外で頻発するとDB負荷増"),
        ("①", "①DB接続プール", "cli_datasource.csv", "物理接続(区間)", "アイドル破棄", "Δ TimedOut", "本", "区間増分", "TimedOut - 前回値", "idle-timeout-minutes を経過して破棄された接続数", "試験中に頻発する場合はアイドル時間設定要確認"),
        ("①", "①DB接続プール", "cli_datasource.csv", "XA(区間)", "ロールバック", "Δ XARollbackCount", "件", "区間増分", "XARollbackCount - 前回値", "2フェーズコミット (2PC) のトランザクションでロールバックされた件数", "▲ 1件以上でトランザクション失敗調査要"),

        # --- ① HTTP (Undertow) ---
        ("①", "①HTTP", "cli_undertow.csv", "基本", "リスナー", "listener", "―", "識別子", "server / type / listener", "Undertow の HTTP/HTTPS リスナー識別名 (default-server/http-listener/default 等)", "―"),
        ("①", "①HTTP", "cli_undertow.csv", "リクエスト(区間)", "件数", "Δ request-count", "件", "区間増分", "request-count - 前回値", "区間に該当リスナーが受け付けた HTTP リクエストの総数", "JMeter の送信リクエスト数と整合性を確認"),
        ("①", "①HTTP", "cli_undertow.csv", "リクエスト(区間)", "毎秒", "RPS", "件/秒", "レート", "Δ request-count ÷ 経過秒", "区間内の秒間平均リクエスト処理件数 (Throughput)", "試験の目標スループット (RPS) 到達確認"),
        ("①", "①HTTP", "cli_undertow.csv", "リクエスト(区間)", "エラー", "Δ error-count", "件", "区間増分", "error-count - 前回値", "500 番台 (Internal Server Error) を返したリクエスト数", "■ 1件以上で要対応。アプリ障害の直接証跡"),
        ("①", "①HTTP", "cli_undertow.csv", "リクエスト(区間)", "エラー率", "エラー率", "%", "比率", "Δ error-count ÷ Δ request-count", "リクエストに対するエラー発生割合", "▲ 0% 超で警告、1% 超で重大"),
        ("①", "①HTTP", "cli_undertow.csv", "処理時間(ms)", "平均(区間)", "avg-processing-time", "ms", "区間平均", "Δ processing-time ÷ Δ request-count", "Undertow 内部でリクエストを処理した平均時間 (record-request-start-time=true 時のみ)", "JMeter側の応答時間との差分でNW遅延特定"),

        # --- ① セッション ---
        ("①", "①セッション", "cli_session.csv", "セッション数", "アクティブ", "active-sessions", "個", "生値", "active-sessions", "現在メモリ上に保持されている HTTP セッションオブジェクト数", "セッションオブジェクト増殖によるヒープ圧迫監視"),
        ("①", "①セッション", "cli_session.csv", "区間", "作成 毎秒", "作成レート", "個/秒", "レート", "Δ sessions-created ÷ 経過秒", "1秒あたりに新規生成されたセッション数", "JMeterシナリオのログイン頻度と一致するか"),
        ("①", "①セッション", "cli_session.csv", "区間", "拒否", "Δ rejected-sessions", "個", "区間増分", "rejected-sessions - 前回値", "max-active-sessions 上限に達してセッション作成が拒否された回数", "■ 1回以上で要対応。セッション数上限設定見直し"),

        # --- ① スレッドプール (EJB / EE) ---
        ("①", "①スレッドプール", "cli_threadpool.csv", "スレッド", "実行中", "active", "本", "生値", "active", "現在タスクを実行しているワーカースレッド数", "プール上限に対する逼迫度合い確認"),
        ("①", "①スレッドプール", "cli_threadpool.csv", "待ち", "キュー", "queue", "件", "生値", "queue-size", "スレッド空き待ちのためタスクキューに滞留しているタスク数", "▲ 1件以上で処理遅延発生中"),
        ("①", "①スレッドプール", "cli_threadpool.csv", "区間", "拒否", "Δ rejected", "件", "区間増分", "rejected - 前回値", "キューが溢れてタスクの実行が拒否された回数", "■ 1件以上で要対応。タスクドロップ発生"),

        # --- ① IO ワーカー ---
        ("①", "①IOワーカー", "cli_io_worker.csv", "タスクスレッド", "使用中", "busy-task-thread-count", "本", "生値", "busy-task-thread-count", "HTTPリクエストを処理しているワーカースレッド数 (default task-N)", "上限に近づいていないか"),
        ("①", "①IOワーカー", "cli_io_worker.csv", "タスクスレッド", "上限", "max-pool-size", "本", "設定", "max-pool-size (CPU数×16が既定)", "タスクスレッドの最大同時起動数上限", "④の判定でCPU数が過大だとこの値も過大になる"),
        ("①", "①IOワーカー", "cli_io_worker.csv", "タスクスレッド", "使用率", "使用中÷上限", "%", "比率", "busy ÷ max-pool-size", "タスクスレッドプールの飽和度合", "▲ 90% 以上で注意、100% でリクエスト待ち"),
        ("①", "①IOワーカー", "cli_io_worker.csv", "タスクスレッド", "キュー", "queue-size", "件", "生値", "queue-size", "ワーカースレッド空き待ちで待機しているリクエスト数", "▲ 1件以上でリクエスト待機発生"),
        ("①", "①IOワーカー", "cli_io_worker.csv", "I/O スレッド", "数", "io-thread-count", "本", "設定", "io-thread-count (CPU数×2が既定)", "ノンブロッキング NIO ネットワーク読み書きを担当するスレッド数", "CPU数に基づいて自動決定される"),

        # --- ① トランザクション ---
        ("①", "①トランザクション", "cli_transactions.csv", "区間", "コミット 毎秒", "TPS", "件/秒", "レート", "Δ committed ÷ 経過秒", "1秒あたりに正常完了・コミットされたトランザクション数", "DB更新処理の総合スループット指標"),
        ("①", "①トランザクション", "cli_transactions.csv", "区間", "中断", "Δ aborted", "件", "区間増分", "aborted - 前回値", "ロールバック等により中断されたトランザクション数", "▲ 増加時はアプリ例外やデッドロック調査"),
        ("①", "①トランザクション", "cli_transactions.csv", "区間", "タイムアウト", "Δ timed-out", "件", "区間増分", "timed-out - 前回値", "トランザクションタイムアウトを超過して強制ロールバックされた件数", "■ 1件以上で要対応。処理の重いSQLやロック待ち"),

        # --- ① JVM (CLI) ---
        ("①", "①JVM（CLI）", "cli_jvm.csv", "取得", "応答 (1=成功)", "cli_ok", "フラグ", "状態", "cli_ok (1:正常, 0:無応答)", "JBoss 管理窓口 (jboss-cli) が正常に応答したか (高負荷や長大GC停止で0になる)", "■ 0 の回がある場合は JBoss が一時無応答"),
        ("①", "①JVM（CLI）", "cli_jvm.csv", "スレッド", "ライブ", "live_threads", "本", "生値", "live_threads", "JVM 全体で稼働しているアクティブな Java スレッドの総数", "スレッドリークの有無を確認"),
        ("①", "①JVM（CLI）", "cli_jvm.csv", "ヒープ(MB)", "使用率", "heap_used÷heap_max", "%", "比率", "heap_used ÷ heap_max", "最大ヒープに対する現在のヒープ使用率", "▲ 85% 以上でGC頻発、メモリ不足の疑い"),

        # --- ② ヒープ・GC (jstat) ---
        ("②", "②ヒープ・GC（jstat）", "jvm_jstat_gc.csv", "ヒープ(MB)", "Eden 使用", "EU", "MB", "生値", "EU ÷ 1024", "Eden領域 (短命オブジェクト生成領域) の現在使用量", "Young GC のトリガーとなる領域"),
        ("②", "②ヒープ・GC（jstat）", "jvm_jstat_gc.csv", "ヒープ(MB)", "Survivor 使用", "S0U+S1U", "MB", "生値", "(S0U + S1U) ÷ 1024", "Survivor領域 (Young GCを生き残ったオブジェクト領域) の使用量", "急増時はオブジェクトの早期昇格が発生"),
        ("②", "②ヒープ・GC（jstat）", "jvm_jstat_gc.csv", "ヒープ(MB)", "Old 使用", "OU", "MB", "生値", "OU ÷ 1024", "Old領域 (長命オブジェクト領域) の現在使用量", "メモリリークやセッション蓄積で右肩上がりに"),
        ("②", "②ヒープ・GC（jstat）", "jvm_jstat_gc.csv", "ヒープ(MB)", "Old ÷ 最大ヒープ", "OU÷MaxHeap", "%", "比率", "OU ÷ MaxHeapMB", "最大ヒープ容量に対する Old 領域使用量の実効占有割合", "▲ 85% 以上で Full GC / メモリ枯渇の危機"),
        ("②", "②ヒープ・GC（jstat）", "jvm_jstat_gc.csv", "GC(区間)", "Young 回数", "Δ YGC", "回", "区間増分", "YGC - 前回値", "区間内に発生した Young GC (Minor GC) の回数", "頻度が高すぎる場合はオブジェクト生成過多"),
        ("②", "②ヒープ・GC（jstat）", "jvm_jstat_gc.csv", "GC(区間)", "Young 時間 ms", "Δ YGCT", "ms", "区間増分", "(YGCT - 前回値) × 1000", "区間内の Young GC による総停止時間", "スループット低下への影響評価"),
        ("②", "②ヒープ・GC（jstat）", "jvm_jstat_gc.csv", "GC(区間)", "Full 回数", "Δ FGC", "回", "区間増分", "FGC - 前回値", "区間内に発生した Full GC の回数 (Stop The World で全停止)", "■ 1回でも発生していれば要対応 (G1チューニング要)"),
        ("②", "②ヒープ・GC（jstat）", "jvm_jstat_gc.csv", "GC(区間)", "GC 時間率", "Δ GCT÷経過", "%", "比率", "Δ GCT ÷ 経過秒", "経過時間のうち GC の停止時間が占めた割合 (スループット損失率)", "▲ 5% 以上でスループット悪化、10% 超で重大"),
        ("②", "②ヒープ・GC（jstat）", "jvm_jstat_gc.csv", "GC(区間)", "最後の GC の原因", "LGCC", "―", "状態", "LGCC (G1 Evacuation Pause 等)", "直近で実行された GC のトリガー要因 (Allocation Failure, Metadata GC Threshold 等)", "原因に応じたヒープ/メタスペース設定見直し"),
        ("②", "②ヒープ・GC（jstat）", "jvm_jstat_gc.csv", "セーフポイント(区間)", "停止 ms", "Δ SPT", "ms", "区間増分", "(SPT - 前回値) × 1000", "GC や JIT デオプティマイズ等のため全スレッドを停止させた合計時間", "GC時間との乖離が大きい場合JVMセーフポイント調査"),

        # --- ② ヒープ (jcmd) ---
        ("②", "②ヒープ（jcmd）", "jvm_jcmd_heap.csv", "G1 リージョン", "サイズ(KB)", "region_size_kb", "KB", "設定", "region_size_kb", "G1 GC がヒープを分割管理するリージョン1つのサイズ (1MB〜32MB)", "巨大オブジェクト (Humongous) の閾値に関与"),
        ("②", "②ヒープ（jcmd）", "jvm_jcmd_heap.csv", "G1 リージョン", "Young 数", "young_regions", "個", "生値", "young_regions", "現在 Young 世代として割り当てられているリージョンの個数", "G1 の自動伸縮挙動の確認"),
        ("②", "②ヒープ（jcmd）", "jvm_jcmd_heap.csv", "メタスペース(MB)", "使用", "metaspace_used_kb", "MB", "生値", "metaspace_used_kb ÷ 1024", "クラスメタデータが格納される Metaspace の実使用量", "クラスの動的ロード増加で肥大化していないか"),

        # --- ② コードキャッシュ ---
        ("②", "②コードキャッシュ", "jvm_codecache.csv", "容量(MB)", "使用率", "使用÷サイズ", "%", "比率", "used_kb ÷ size_kb", "JIT コンパイル済みネイティブコード置き場の使用率", "▲ 90% 以上で空き枯渇警告"),
        ("②", "②コードキャッシュ", "jvm_codecache.csv", "コンパイル済み", "満杯の回数", "full_count", "回", "生値", "full_count", "コードキャッシュが満杯になった回数。0以外になるとJITコンパイラが停止し深刻な速度低下", "■ 1回以上で要対応 (-XX:ReservedCodeCacheSize 拡大要)"),

        # --- ② ネイティブメモリ (NMT) ---
        ("②", "②NMT", "jvm_nmt.csv", "合計(MB)", "コミット", "Total committed", "MB", "生値", "Total.committed_kb ÷ 1024", "JVM 全体が OS からコミット (物理メモリ確保) している総メモリ量", "コンテナのメモリ制限 (RAM Limit) との比較"),
        ("②", "②NMT", "jvm_nmt.csv", "分類別(MB)", "Java Heap", "Java Heap", "MB", "生値", "Java Heap.committed_kb ÷ 1024", "Java ヒープ領域の実コミット量", "最大ヒープ設定 (-Xmx) に応じて確保"),
        ("②", "②NMT", "jvm_nmt.csv", "分類別(MB)", "Class / Thread / GC", "Class / Thread / GC", "MB", "生値", "各カテゴリの committed_kb ÷ 1024", "メタスペース、スレッドスタック (スレッド数×-Xss)、GC管理テーブル等のメモリ内訳", "ヒープ外メモリ増大 (オフヒープリーク) の特定"),

        # --- ② スレッド状態 (jcmd Thread.print) ---
        ("②", "②スレッド状態", "jvm_thread_states.csv", "全体", "ブロック", "all.blocked", "本", "生値", "all.blocked", "Java ロック (synchronized / モニターロック) の取得待ちでブロックしているスレッド数", "▲ 1本以上でロック競合・ボトルネック発生中"),
        ("②", "②スレッド状態", "jvm_thread_states.csv", "DB 接続待ち", "スレッド", "db-pool-wait.total", "本", "生値", "db-pool-wait.total", "IronJacamar の接続プールで空き接続を待って停滞しているスレッド数", "■ 1本以上で要対応。DB接続プール枯渇の直接証拠"),
        ("②", "②スレッド状態", "jvm_thread_states.csv", "HTTP処理", "処理中", "http-task.busy", "本", "生値", "http-task.busy", "HTTP リクエストを現に処理実行しているタスクスレッド数", "実行スレッド数の推移と上限の余裕度確認"),

        # --- ② GC (JFR) ---
        ("②", "②GC（JFR）", "jfr_gc.csv", "停止(ms)", "最長", "longest_pause_ms", "ms", "生値", "longest_pause_ms", "該当 GC 1回の中でアプリケーションスレッドが連続停止した最長時間", "▲ 200ms 以上で注意 (G1既定停止目標超過)"),
        ("②", "②GC（JFR）", "jfr_gc.csv", "停止(ms)", "合計", "sum_of_pauses_ms", "ms", "生値", "sum_of_pauses_ms", "該当 GC のフェーズ全体での累計一時停止時間", "1回のGCによる合計レイテンシ影響"),
        ("②", "②GC（JFR）", "jfr_gc.csv", "ヒープ(MB)", "回収", "GC前−GC後", "MB", "差分", "heap_before_mb - heap_after_mb", "該当 GC によって回収・解放されたメモリ量", "回収効率が悪い場合はメモリ逼迫の兆候"),

        # --- ② JVM統計 (JFR) ---
        ("②", "②JVM統計（JFR）", "jfr_stats.csv", "CPU(%)", "JVM 合計", "user + system", "%", "比率", "jvm_user_pct + jvm_system_pct", "JVM プロセスが使用した CPU 時間の割合 (JVMから見える全CPU=100%)", "JVM全体の処理負荷水準"),
        ("②", "②JVM統計（JFR）", "jfr_stats.csv", "例外", "毎秒", "throwables/sec", "件/秒", "レート", "Δ throwables ÷ 経過秒", "区間内に発生した Java 例外 (Throwable) の1秒あたり生成件数", "▲ 100件/秒以上で注意 (例外生成スタックトレース負荷)"),
        ("②", "②JVM統計（JFR）", "jfr_stats.csv", "CPU 制限(区間)", "スロットリング回数", "Δ throttled_slices", "回", "区間増分", "throttled_slices - 前回値", "cgroup の CPU CFS クォータ超過により実行を止められた回数", "▲ 1回以上でCPU制限によるスローダウン発生"),

        # --- ② JFR 上位 ---
        ("②", "②JFR上位", "jfr_hot_methods.csv", "Hot Methods", "メソッド名", "method", "―", "集計", "CPU実行プロファイル上位メソッド", "JFR サンプリングで最も CPU 時間を消費していた Java メソッド", "アプリのボトルネックメソッド特定とチューニング"),
        ("②", "②JFR上位", "jfr_contention.csv", "ロック競合", "クラス / 待機箇所", "site", "―", "集計", "ロック待ち時間上位のクラスおよびソースコード位置", "スレッドが最も長く待たされた同期オブジェクトと待機時間", "排他制御 (synchronized / Lock) の見直し箇所"),

        # --- ③ top 概要 ---
        ("③", "③top概要", "os_top_summary.csv", "ロードアベレージ", "1 分", "load1", "値", "生値", "load1", "直近1分間の実行待ちプロセスの平均数 (CPU数との比較で過負荷判定)", "CPU 数を超えている場合は CPU 資源不足"),
        ("③", "③top概要", "os_top_summary.csv", "CPU(%,全CPU)", "us / sy / wa / st", "cpu_us, sy, wa, st", "%", "比率", "top CPU内訳", "VM 全体の CPU 時間内訳 (us:アプリ, sy:カーネル, wa:IO待ち, st:steal)", "▲ st 10% 以上でホスト競合、wa 10% 以上でEFS/ディスク遅延"),
        ("③", "③top概要", "os_top_summary.csv", "プロセス別CPU", "収集処理", "cpu.collect", "%", "集計", "jboss-cli + jstat + jcmd + perf-metrics の %CPU", "本レポート収集処理自身が消費した CPU 使用率 (1コア=100%)", "平均 1.7% 前後であることを確認しオーバーヘッド検証"),

        # --- ③ top スレッド (JBoss CPU 上位) ---
        ("③", "③topスレッド", "os_top_threads.csv", "基本", "スレッド名 / 用途", "thread, tgroup", "―", "識別", "top -H による TID とスレッド名 (用途分類)", "CPU を大量消費している JBoss のスレッド名と用途 (HTTP/GC/JIT/VM)", "スレッドダンプの nid (16進数) と TID (10進数) で突合可能"),
        ("③", "③topスレッド", "os_top_threads.csv", "使用量", "CPU%", "cpu_pct", "%", "生値", "cpu_pct (1コア=100%)", "該当スレッドの CPU 使用率", "▲ 50% 以上の高負荷スレッドの特定"),

        # --- ③ JBoss スレッド (用途別) ---
        ("③", "③JBossスレッド", "os_jboss_threads.csv", "CPU(コア,用途別)", "HTTP処理", "http-task.cpu_cores", "コア", "集計", "http-task の cpu_pct ÷ 100", "Undertow の HTTP リクエスト処理スレッド群が消費した CPU コア数", "アプリ本来の処理に正しく CPU が使われているか"),
        ("③", "③JBossスレッド", "os_jboss_threads.csv", "CPU(コア,用途別)", "GC", "gc.cpu_cores", "コア", "集計", "GC スレッド群が消費した CPU コア数", "▲ JBoss 全体CPUの20%以上をGCが占める場合はメモリ設計見直し"),
        ("③", "③JBossスレッド", "os_jboss_threads.csv", "CPU(コア,用途別)", "JIT", "jit.cpu_cores", "コア", "集計", "JIT コンパイラスレッド (C1/C2) が消費した CPU コア数", "ウォームアップ段階と定常状態での JIT 負荷推移確認"),
        ("③", "③JBossスレッド", "os_jboss_threads.csv", "全体", "非自発的切替/秒", "all.nonvol_ctxsw_per_sec", "回/秒", "生値", "/proc/<pid>/status の nonvoluntary_ctxt_switches", "CPUタイムスライスの途中で強制的にCPUを取り上げられた回数", "多いほど CPU コア数に対してスレッドが過多で競合中"),

        # --- ③ FD・TCP ---
        ("③", "③FD・TCP", "os_jboss_proc.csv", "ファイル記述子", "使用率", "fd_count÷fd_limit", "%", "比率", "fd_count ÷ fd_limit", "JBoss プロセスが保持するファイル記述子 (ソケット+ファイル) の使用率", "■ 80% 以上で要対応。FD枯渇による Too many open files 防災"),
        ("③", "③FD・TCP", "os_jboss_proc.csv", "TCP(JBoss)", "CLOSE_WAIT", "tcp_close_wait", "本", "生値", "tcp_close_wait", "相手から切断されたのに JBoss 側で close() していないソケット数", "▲ 1本以上で増加傾向なら接続クローズ漏れ (リーク)"),
        ("③", "③FD・TCP", "os_jboss_proc.csv", "TCP(タスク全体)", "TIME_WAIT", "tcp_time_wait_all", "本", "生値", "tcp_time_wait_all", "タスク全体で TIME_WAIT 状態にあるソケット数", "▲ 1000本以上で注意。HTTP Keep-Alive 接続再利用の不備"),

        # --- ④ CPU 判定 & スケジューリング ---
        ("④", "④CPU判定", "cpu_env.csv", "判定結果", "判定", "judgement.result", "判定", "状態", "要 / 要検討 / 不要 / 設定済み / 判定不可", "Fargate の CPU シェア配分と JVM 認識 CPU 数の乖離に基づく判定結果", "「要」の場合は -XX:ActiveProcessorCount の設定が必須"),
        ("④", "④CPU判定", "cpu_env.csv", "判定結果", "推奨設定値", "recommended_val", "設定", "算出値", "-XX:ActiveProcessorCount=<コンテナ取り分vCPU (四捨五入・最小1)>", "JBoss の JVM 起動オプション (PREPEND_JAVA_OPTS) に足すべき推奨パラメータ", "指定により過剰なGC/JIT/Undertowスレッド生成を抑制"),
        ("④", "④CPUスケジューリング", "cpu_sched.csv", "JBossのCPU(コア)", "合計 vs 取り分", "jboss_cpu_cores, share_vcpu", "コア", "比較", "JBoss実測CPU使用量 と タスク定義上のコンテナ取り分 (vCPU)", "取り分を超過してCPUを使えているか、あるいは取り合いで制限されているか", "▲ 他コンテナ同居時に取り分を超えていると将来遅延"),
        ("④", "④CPUスケジューリング", "cpu_sched.csv", "CPUの待ち(PSI)", "some avg10 (%)", "psi_some_avg10", "%", "生値", "/proc/pressure/cpu some avg10", "Linux カーネルの PSI (CPU 資源の不足によりプロセスが待たされた時間の割合)", "▲ 95%値で 20% 以上なら深刻な CPU 不足"),
        ("④", "④CPUスケジューリング", "cpu_sched.csv", "実行待ち", "実行可能(全体)", "procs_running", "スレッド", "生値", "/proc/stat の procs_running", "実行中または CPU 実行待ち行列に並んでいるスレッドの総数", "CPU 数 (OS) を超えている時間帯は CPU 待ちが発生")
    ]

    r_idx = 4
    for row_data in metrics_data:
        ws2.row_dimensions[r_idx].height = 28 if len(row_data[9]) > 30 else 22
        fill_cur = fill_row_even if r_idx % 2 == 0 else fill_row_odd
        k = row_data[0]
        
        c0 = ws2.cell(row=r_idx, column=1, value=k)
        c0.font = f_badge; c0.fill = fill_badges.get(k, fill_th); c0.alignment = al_center; c0.border = bd_thin
        
        for ci, val in enumerate(row_data[1:], start=2):
            cell = ws2.cell(row=r_idx, column=ci, value=val)
            cell.font = f_data; cell.fill = fill_cur; cell.border = bd_thin
            if ci in [2, 7, 8]:
                cell.alignment = al_center
            else:
                cell.alignment = al_left
        r_idx += 1

    ws2.auto_filter.ref = f"A3:K{r_idx - 1}"

    # =============================================================
    # SHEET 3: ②実装コマンドリファレンス (Implementation Commands)
    # =============================================================
    ws3 = wb.create_sheet(title="②実装コマンドリファレンス")
    apply_title_banner(ws3, "メトリクス取得・実装コマンド完全リファレンス", 
                       "現場でそのままコピー＆ペーストして実行・実装できる Bash / jboss-cli / jstat / jcmd / JFR / proc / AWS CLI コマンド集", 8)
    ws3.views.sheetView[0].showGridLines = True
    ws3.freeze_panes = "D4"
    ws3.row_dimensions[3].height = 26

    cols_sh3 = [
        ("区分", 8, al_center),
        ("ツール / サブシステム", 18, al_center),
        ("取得メトリクス / 目的", 24, al_left),
        ("実行環境 / 権限", 18, al_center),
        ("そのまま実装・利用できるコマンド (ワンライナー / 実装例)", 68, al_cmd),
        ("コマンドの詳細解説・主要引数", 36, al_left),
        ("実装上の注意点・前提条件 (ボリューム/パーミッション等)", 36, al_left),
        ("レポート出力先ファイル", 20, al_left)
    ]

    for ci, (cn, cw, cal) in enumerate(cols_sh3, start=1):
        cell = ws3.cell(row=3, column=ci, value=cn)
        cell.font = f_th; cell.fill = fill_th; cell.alignment = al_center; cell.border = bd_th
        ws3.column_dimensions[get_column_letter(ci)].width = cw

    # 実装コマンドデータリスト
    commands_data = [
        # --- ① jboss-cli ---
        ("①", "jboss-cli", "DB接続プール統計の即時取得\n(全データソース一括)", "コンテナ内 (sh / root or jboss)\nまたは ECS Exec",
         '${JBOSS_HOME:-/opt/jboss-eap}/bin/jboss-cli.sh --connect --controller=remote+http://127.0.0.1:9990 \\\n  --command="/subsystem=datasources/data-source=*:read-resource(include-runtime=true)"',
         "ワイルドカード (*) で全データソースの ActiveCount, InUseCount, WaitCount, BlockingFailureCount 等の DMR を一度に取得。",
         "JBoss EAP 起動完了後 (ポート9990リスン後) に実行可能。ローカル接続では認証不要。", "cli_datasource.csv"),

        ("①", "jboss-cli", "XAデータソース接続プール統計の即時取得", "コンテナ内 / ECS Exec",
         '${JBOSS_HOME:-/opt/jboss-eap}/bin/jboss-cli.sh --connect --controller=remote+http://127.0.0.1:9990 \\\n  --command="/subsystem=datasources/xa-data-source=*:read-resource(include-runtime=true)"',
         "XA データソースの接続プール統計および XACommitCount, XARollbackCount, XAPrepareCount 等を取得。",
         "2フェーズコミットを使用する分散トランザクション対象データソース用。", "cli_datasource.csv"),

        ("①", "jboss-cli", "プール統計が未有効な場合の\n動的有効化コマンド", "コンテナ内 / ECS Exec",
         '${JBOSS_HOME:-/opt/jboss-eap}/bin/jboss-cli.sh --connect --controller=remote+http://127.0.0.1:9990 \\\n  --command="/subsystem=datasources/data-source=<DS_NAME>:write-attribute(name=statistics-enabled,value=true)"',
         "統計が無効 (false) だとメトリクス値が 0 のままになるため、JBoss 稼働中に動的に統計カウンタを ON にする。",
         "設定ファイル (standalone.xml) は変更されず、再起動なしで即座に有効化される。", "cli_datasource.csv"),

        ("①", "jboss-cli", "HTTP リスナー (Undertow) 統計の取得", "コンテナ内 / ECS Exec",
         '${JBOSS_HOME:-/opt/jboss-eap}/bin/jboss-cli.sh --connect --controller=remote+http://127.0.0.1:9990 \\\n  --command="/subsystem=undertow/server=default-server/http-listener=*:read-resource(include-runtime=true)"',
         "Undertow の request-count, error-count (500系エラー), processing-time などを取得。",
         "平均処理時間を取得するには listener の record-request-start-time=true が必要。", "cli_undertow.csv"),

        ("①", "jboss-cli", "セッション統計の取得\n(デプロイメント別)", "コンテナ内 / ECS Exec",
         '${JBOSS_HOME:-/opt/jboss-eap}/bin/jboss-cli.sh --connect --controller=remote+http://127.0.0.1:9990 \\\n  --command="/deployment=*/subsystem=undertow:read-resource(include-runtime=true)"',
         "デプロイされている WAR アプリごとの active-sessions, sessions-created, rejected-sessions を取得。",
         "セッション上限による拒否 (rejected) や急激なセッション増加を監視可能。", "cli_session.csv"),

        ("①", "jboss-cli", "IO ワーカー (スレッド数・キュー) の取得", "コンテナ内 / ECS Exec",
         '${JBOSS_HOME:-/opt/jboss-eap}/bin/jboss-cli.sh --connect --controller=remote+http://127.0.0.1:9990 \\\n  --command="/subsystem=io/worker=*:read-resource(include-runtime=true)"',
         "Undertow の HTTP 処理スレッド (busy-task-thread-count, max-pool-size, queue-size, io-thread-count) を取得。",
         "リクエスト処理スレッドの枯渇や待機キュー発生の検出に必須。", "cli_io_worker.csv"),

        ("①", "jboss-cli", "トランザクション統計の取得", "コンテナ内 / ECS Exec",
         '${JBOSS_HOME:-/opt/jboss-eap}/bin/jboss-cli.sh --connect --controller=remote+http://127.0.0.1:9990 \\\n  --command="/subsystem=transactions:read-resource(include-runtime=true)"',
         "Narayana トランザクションマネージャーの committed, aborted, timed-out, inflight 数を取得。",
         "起動時引数 -Dwildfly.transactions.statistics-enabled=true が必要。", "cli_transactions.csv"),

        ("①", "jboss-cli", "JVM プラットフォーム MBean 情報の取得", "コンテナ内 / ECS Exec",
         '${JBOSS_HOME:-/opt/jboss-eap}/bin/jboss-cli.sh --connect --controller=remote+http://127.0.0.1:9990 \\\n  --command="/core-service=platform-mbean/type=threading:read-resource(include-runtime=true)"',
         "JVM 内の live_threads, daemon_threads, peak_threads などのスレッド状態を取得。",
         "memory, operating-system, class-loading, buffer-pool/name=direct 等も同様に取得可能。", "cli_jvm.csv"),

        ("①", "jboss-cli", "定周期高速収集のための\n常駐 FIFO パイプ実装方式", "シェルスクリプト\n(perf-metrics.sh 内装)",
         'mkfifo /tmp/cli.in /tmp/cli.out\n"${JBOSS_HOME}/bin/jboss-cli.sh" --connect --controller="${CTRL}" < /tmp/cli.in > /tmp/cli.out &\n# パイプ経由でクエリを送り、DMR応答をawkでパース',
         "毎回の CLI 起動 JVM (200MB/数秒) を排し、1 つの常駐プロセスに FIFO 経由でクエリを連続投入。",
         "CPU負荷を劇的に低減。タイムアウト監視 (kill) 機構を併設して高負荷時のハングを防止。", "全 cli_*.csv"),

        # --- ② jstat / jcmd / JFR ---
        ("②", "jstat", "ヒープ領域・GC統計の定周期取得\n(OpenJDK 21 確実実行コマンド)", "コンテナ内 (sh / root or jboss)",
         'java -m jdk.jcmd/sun.tools.jstat.Jstat -gc <JBOSS_PID> 10000',
         "Red Hat OpenJDK runtime イメージに `jstat` バイナリが無い場合でも、JDKモジュール直接指定で10秒ごとに永続出力。",
         "/tmp に書き込み可能なボリュームが必要 (hsperfdata_<user>/<pid> のソケットを読み取るため)。", "jvm_jstat_gc.csv"),

        ("②", "jstat", "カスタムオプション perfgc による\nヒープ・GC・JIT・直近原因の一括取得", "コンテナ内 (sh / root or jboss)",
         'mkdir -p ~/.jvmstat\ncat << \'EOF\' > ~/.jvmstat/jstat_options\noption perfgc {\n  column { header "^S0C^" data sun.gc.generation.0.space.1.capacity / 1024 }\n  column { header "^OU^"  data sun.gc.generation.1.space.0.used / 1024 }\n  column { header "^YGC^" data sun.gc.collector.0.invocations }\n  column { header "^FGC^" data sun.gc.collector.1.invocations }\n  column { header "^LGCC^" data sun.gc.lastCause }\n}\nEOF\njava -m jdk.jcmd/sun.tools.jstat.Jstat -perfgc <JBOSS_PID> 10000',
         "~/.jvmstat/jstat_options を定義することで、-gc, -gccause, -compiler, -safepoint の全項目を1プロセスで同時採取可能。",
         "perf-metrics.sh が自動生成して使用。GC停止時間・セーフポイント停止時間・JITコンパイル数がすべて揃う。", "jvm_jstat_gc.csv"),

        ("②", "jcmd", "ヒープ詳細・G1リージョン情報の取得", "コンテナ内 / ECS Exec",
         'jcmd <JBOSS_PID> GC.heap_info',
         "G1 GC のリージョンサイズ (region size)、Young/Survivor リージョン数、Metaspace/Class space のコミット量を正確に出力。",
         "短時間の JVM 起動を伴うため、本スクリプトでは 60 秒間隔で実行。", "jvm_jcmd_heap.csv"),

        ("②", "jcmd", "JIT コードキャッシュ使用状況の取得", "コンテナ内 / ECS Exec",
         'jcmd <JBOSS_PID> Compiler.codecache',
         "non-nmethods, profiled nmethods, non-profiled nmethods のサイズ・使用量・full_count を出力。",
         "full_count > 0 は JIT が停止したことを意味するため即時対処が必要。", "jvm_codecache.csv"),

        ("②", "jcmd", "スレッド一覧・状態・ロック待ちの取得\n(Thread.print)", "コンテナ内 / ECS Exec",
         'jcmd <JBOSS_PID> Thread.print',
         "全 Java スレッドのスタックトレースと状態を出力。awk で http-task, db-pool-wait, blocked などの用途別に即時集計可能。",
         "「DB接続待ち (IronJacamar)」や「synchronized ロック待ち」のスレッド本数を特定する。", "jvm_thread_states.csv"),

        ("②", "jcmd", "ネイティブメモリ追跡 (NMT) の取得", "コンテナ内 / ECS Exec",
         'jcmd <JBOSS_PID> VM.native_memory summary',
         "Java Heap, Class, Thread, Code, GC, Compiler 等のネイティブメモリ予約量・コミット量を出力。",
         "起動オプション `-XX:NativeMemoryTracking=summary` が必要 (PERF_NMT=summary で有効化)。", "jvm_nmt.csv"),

        ("②", "jcmd", "JVM バージョン・フラグ・起動引数の取得", "コンテナ内 / ECS Exec",
         'jcmd <JBOSS_PID> VM.version\njcmd <JBOSS_PID> VM.command_line\njcmd <JBOSS_PID> VM.flags -all',
         "JVMの完全ビルド情報、設定された起動引数、すべてのエルゴノミクス適用済み内部フラグを網羅ダンプ。",
         "ActiveProcessorCount や ParallelGCThreads の実効値を確定させるために使用。", "jvm_info.txt\ncpu_env.csv"),

        ("②", "JFR", "JBoss 起動時からの低負荷 JFR 記録\n(JVM 起動引数への組み込み)", "タスク定義環境変数\nPREPEND_JAVA_OPTS",
         '-XX:StartFlightRecording=dumponexit=true,filename=/mnt/logs/.../perf/jfr/perf_%p_%t.jfr,settings=heap,maxage=24h,maxsize=1g',
         "JBoss 起動時から停止時まで Flight Recorder をバックグラウンド記録。settings=heap でヒープ・GC・統計に限定し低負荷化。",
         "PREPEND_JAVA_OPTS で渡すことで standalone.conf の既定メモリ設定を破壊せずに適用。", "jfr/*.jfr"),

        ("②", "JFR", "試験終了時の手動 JFR ダンプ実行", "コンテナ内 / ECS Exec",
         'jcmd <JBOSS_PID> JFR.dump name=perf filename=/mnt/logs/.../perf/jfr/manual_dump.jfr',
         "JBoss を停止させずに、記録中の JFR バッファをファイルへフラッシュ・書き出し。",
         "perf-metrics.sh jfr-dump コマンドで一発実行可能。", "jfr/dump/*.jfr"),

        ("②", "JFR", "単一 Java ソースによる\nJFR から CSV への高速自動抽出", "コンテナ内 または PC\n(JDK 17/21 単独実行)",
         'java test/fixtures/perf_sample/jfr/PerfJfrCsv.java <input.jfr> <output_dir> 10',
         "JDK 標準の jdk.jfr.consumer API を使用し、GCごとの停止時間・Hot Methods・ロック競合・IOイベントを CSV 群に変換。",
         "外部ライブラリ一切不要。単一ソース直接実行 (java X.java)。perf-metrics.sh jfr-csv または レポートの -ConvertJfr で呼出。", "jfr_*.csv"),

        # --- ③ top / ps / /proc ---
        ("③", "top", "システム概要・プロセス別 CPU の取得\n(バッチモード)", "コンテナ内 (sh / root or jboss)",
         'top -b -n 1',
         "ロードアベレージ、CPU 内訳 (us, sy, id, wa, st)、タスク数、プロセス一覧の CPU/メモリ使用率を1回採取。",
         "procps-ng がインストールされている場合に使用。-b オプション必須。", "os_top_summary.csv\nos_top_process.csv"),

        ("③", "top", "JBoss スレッド別 CPU 上位 10 件の取得", "コンテナ内 (sh / root or jboss)",
         'top -H -b -n 1 -p <JBOSS_PID> | head -n 17',
         "スレッド単位 (-H) で CPU 使用率の高い順にソートして出力。TID (スレッドID) と CPU% を取得。",
         "どのスレッド (C2コンパイラ, GCスレッド, ワーカースレッド等) が CPU を喰っているか特定。", "os_top_threads.csv"),

        ("③", "ps", "全プロセスの詳細状態・スレッド数の取得", "コンテナ内 (sh / root or jboss)",
         'ps -eo pid,ppid,user,stat,nlwp,rss,vsz,pcpu,pmem,etimes,cputime,args --sort=-pcpu',
         "プロセスごとのスレッド数 (NLWP)、常駐メモリ (RSS)、仮想メモリ (VSZ)、起動来CPU%、経過時間を出力。",
         "収集処理やサブプロセスがゾンビ化していないか、想定外のスレッド急増がないかを監視。", "os_ps.csv"),

        ("③", "/proc", "procps-ng 不在環境での\n/proc からの JBoss スレッド用途別集計", "コンテナ内 (sh / root or jboss)",
         'for f in /proc/<JBOSS_PID>/task/*/stat; do\n  # スレッド名 (comm) と CPU tick (utime+stime) を読み取り、用途別に awk で集計\n  cat "$f"\ndone',
         "UBI Minimal 等で top が使えない場合でも、カーネルの /proc/<pid>/task から完全同一の用途別 CPU と実行状態を取得。",
         "perf-metrics.sh が自動フォールバック。http-task, http-io, gc, jit, jfr, vm 等に完全分類。", "os_jboss_threads.csv"),

        ("③", "/proc", "ファイル記述子 (FD) 数と上限の取得", "コンテナ内 (sh / root or jboss)",
         'echo "FD_COUNT: $(ls -1 /proc/<JBOSS_PID>/fd 2>/dev/null | wc -l)"\ngrep "Max open files" /proc/<JBOSS_PID>/limits',
         "JBoss が現在オープンしているファイル・ソケット・パイプの総数と、プロセスの上限 (soft/hard) を取得。",
         "使用率が 80% を超えると Too many open files 例外のリスク大。", "os_jboss_proc.csv"),

        ("③", "/proc", "TCP ソケット状態・ポート別の集計", "コンテナ内 (sh / root or jboss)",
         'cat /proc/net/tcp /proc/net/tcp6 | awk \'\nNR>1 { state=$4; if (state=="01") est++; else if (state=="08") cw++; else if (state=="06") tw++ }\nEND { print "EST="est, "CLOSE_WAIT="cw, "TIME_WAIT="tw }\'',
         "/proc/net/tcp のステートフィールド (01:ESTABLISHED, 08:CLOSE_WAIT, 06:TIME_WAIT, 0A:LISTEN) を直接解析。",
         "netstat や ss が無くてもカーネルから直接接続数を集計。CLOSE_WAIT リークや TIME_WAIT 滞留を検知。", "os_jboss_proc.csv"),

        # --- ④ CPU 判定 & cgroup & ECS ---
        ("④", "ECS Metadata", "ECS タスクメタデータ v4 からの\nコンテナ取り分 CPU 算出", "コンテナ内 (curl / sh)",
         'curl -s "${ECS_CONTAINER_METADATA_URI_V4}/task" > /tmp/task.json\ncurl -s "${ECS_CONTAINER_METADATA_URI_V4}" > /tmp/container.json',
         "タスク定義の CPU 制限 (Limits.CPU) とコンテナ定義の CPU (64ユニット以上) を取得し、コンテナの取り分 vCPU を計算。",
         "Fargate 1.4.0 以降の環境変数 ECS_CONTAINER_METADATA_URI_V4 からローカル取得。", "cpu_env.csv"),

        ("④", "cgroup", "cgroup v2 / v1 CPU クォータ・制限の取得", "コンテナ内 (sh)",
         '# cgroup v2\nif [ -f /sys/fs/cgroup/cpu.max ]; then cat /sys/fs/cgroup/cpu.max; cat /sys/fs/cgroup/cpu.stat; fi\n# cgroup v1\nif [ -f /sys/fs/cgroup/cpu/cpu.cfs_quota_us ]; then cat /sys/fs/cgroup/cpu/cpu.cfs_quota_us; cat /sys/fs/cgroup/cpu/cpu.shares; fi',
         "CFS クォータ (上限) と CFS ピリオド、CPU シェア値、スロットリング発生統計を取得。",
         "Fargate ではクォータは無制限 (max) でシェアのみ配分されるため、JVMが全ホストCPUを認識してしまう。", "cpu_env.csv\ncpu_sched.csv"),

        ("④", "Linux カーネル", "CPU PSI (Pressure Stall Information) の取得", "コンテナ内 (sh)",
         'cat /proc/pressure/cpu',
         "some avg10, avg60, avg300 および累積待機時間 (total) を取得。",
         "CPU 不足によって実行可能なプロセスが待たされた時間の割合をカーネルが直接計測した真の飽和指標。", "cpu_sched.csv"),

        ("④", "JVM プローブ", "JVM が認識した CPU 数の確認と\nActiveProcessorCount 判定", "コンテナ内 / ECS Exec",
         'java -XshowSettings:system -version 2>&1 | grep -E "(Operating System Metrics|active_processor_count)"\n# または判定スクリプト直接実行\n/usr/local/bin/perf-metrics.sh cpu-check',
         "JVM がホストOSやcgroupから自動認識した CPU コア数を調査。コンテナの取り分と照合して「要/不要」を判定。",
         "判定結果「要」なら `-XX:ActiveProcessorCount=<取り分vCPU (四捨五入・最小1)>` の適用を即時リコメンド。", "cpu_env.csv\nサマリー"),

        # --- レポート生成 ---
        ("共通", "PowerShell", "Excel レポートの自動生成\n(単一 perf/ ディレクトリ)", "Windows 11 PowerShell 5.1 / 7",
         'powershell -NoProfile -ExecutionPolicy Bypass -File .\\report\\New-PerfMetricsReport.ps1 \\\n  D:\\perf-results\\perf -OutputPath D:\\reports',
         "perf/ 内の全 CSV を解析し、Meiryo UI モノトーン形式で最大29シート・49グラフの Excel ブック (.xlsx) を生成。",
         "Excel 本体のインストール不要 (OpenXML ZIP を直接構築)。わずか数秒〜1分で完了。", "Excel ブック"),

        ("共通", "PowerShell", "親ディレクトリ一括処理 / 期間指定 /\nJFR 自動変換付きレポート生成", "Windows 11 PowerShell 5.1 / 7",
         '.\\report\\New-PerfMetricsReport.ps1 D:\\efs-copy\\logs -OutputPath D:\\reports \\\n  -From "2026-10-24 23:12" -To "2026-10-24 23:14" -ConvertJfr -Open',
         "配下の全コンテナ (front/back) の perf ディレクトリを自動探索して個別レポート化。JMeter試験時間帯に限定抽出可能。",
         "-ConvertJfr を指定すると PC 側の Java 17/21 で .jfr を自動解析して取り込む。", "複数 Excel ブック")
    ]

    r_idx = 4
    for row_data in commands_data:
        cmd_text = row_data[4]
        desc_text = row_data[5]
        line_count = max(cmd_text.count('\n') + 1, (len(desc_text) // 30) + 1, 2)
        ws3.row_dimensions[r_idx].height = max(line_count * 18, 36)
        
        fill_cur = fill_row_even if r_idx % 2 == 0 else fill_row_odd
        k = row_data[0]

        c0 = ws3.cell(row=r_idx, column=1, value=k)
        c0.font = f_badge; c0.fill = fill_badges.get(k, fill_th); c0.alignment = al_center; c0.border = bd_thin

        for ci, val in enumerate(row_data[1:], start=2):
            cell = ws3.cell(row=r_idx, column=ci, value=val)
            cell.border = bd_thin
            if ci == 5:
                cell.font = f_cmd
                cell.fill = fill_cmd if r_idx % 2 == 0 else fill_cmd_alt
                cell.alignment = al_cmd
            elif ci in [2, 4]:
                cell.font = f_data_bold if ci == 2 else f_data
                cell.fill = fill_cur
                cell.alignment = al_center
            else:
                cell.font = f_data
                cell.fill = fill_cur
                cell.alignment = al_left
        r_idx += 1

    ws3.auto_filter.ref = f"A3:H{r_idx - 1}"

    # =============================================================
    # SHEET 4: ③サマリー所見・49グラフ定義 (Summary Findings & Charts)
    # =============================================================
    ws4 = wb.create_sheet(title="③サマリー所見・49グラフ定義")
    apply_title_banner(ws4, "サマリー所見自動判定ルール & 全 49 グラフ構成定義", 
                       "New-PerfMetricsReport.ps1 がレポート先頭に自動出力する所見の重要度判定ロジックとグラフ一覧", 5)
    ws4.views.sheetView[0].showGridLines = True
    ws4.row_dimensions[3].height = 10

    # 1. 所見ルール
    ws4.cell(row=4, column=1, value="1. サマリーシート「所見」の重要度別自動判定ルール一覧").font = f_section
    merge_and_style(ws4, 4, 1, 4, 5)

    cols_find = ["重要度", "分類", "判定条件 / トリガー指標", "自動出力メッセージ概要", "推奨アクション・恒久対策"]
    r_idx = 5
    ws4.row_dimensions[r_idx].height = 24
    for ci, h in enumerate(cols_find, start=1):
        cell = ws4.cell(row=r_idx, column=ci, value=h)
        cell.font = f_th; cell.fill = fill_th; cell.alignment = al_center; cell.border = bd_th
    r_idx += 1

    findings_rules = [
        ("■ 要対応", "DB 接続プール", "BlockingFailureCount > 0", "「<DS名>」で接続の取得失敗 (取得タイムアウト) が N 件ありました。アプリでは IJ000453 例外になります。", "最大プール数 (max-pool-size) の拡大、またはトランザクション/コネクション保持時間の短縮。"),
        ("■ 要対応", "ヒープ・GC", "Full GC 回数 (FGC) > 0", "Full GC が N 件発生しました。全スレッドが停止 (Stop The World) しました。", "ヒープ容量 (-Xmx) の拡大、Old領域使用率の調査、メモリリークの解消。"),
        ("■ 要対応", "コードキャッシュ", "full_count > 0", "コードキャッシュが満杯になりました。JIT コンパイルが停止し性能が著しく劣化します。", "-XX:ReservedCodeCacheSize を拡大 (例: 256m → 512m)。"),
        ("■ 要対応", "ファイル記述子", "FD 使用率 >= 80%", "ファイル記述子の使用率が最大 N% に達しました (上限 M)。", "DB接続、HTTPクライアント、ファイルストリームのクローズ漏れ (リーク) の解消。"),
        ("■ 要対応", "ActiveProcessorCount", "判定結果 == 「要」", "要 — JVM の CPU 数 (N) > 取り分 (M vCPU)。CPU の取り合いが発生します。", "タスク定義の PREPEND_JAVA_OPTS に `-XX:ActiveProcessorCount=M` を指定。"),
        ("▲ 注意", "DB 接続プール", "使用中 >= 最大プール (AtMax > 0)", "「<DS名>」の使用中が最大プールに達しました (N / M 回)。", "プールの空きバッファがありません。同時アクセス数増加に備えて最大プール枠拡大を検討。"),
        ("▲ 注意", "DB 接続プール", "空き待ち (WaitCount) > 0", "「<DS名>」で空き接続の待ちが N 件発生しました (最大待ち時間 X ms)。", "平均待ちスレッド数を確認し、多重度に見合ったプールサイズへ調整。"),
        ("▲ 注意", "DB 接続待ちスレッド", "db-pool-wait スレッド > 0", "DB 接続の空きを待って止まっているスレッドが最大 N 本ありました (jcmd Thread.print)。", "JBoss アプリ側でコネクション取得待ち行列が発生している直接証拠。"),
        ("▲ 注意", "GC 停止時間", "GC 停止時間率 >= 5%", "GC 時間率が N% でした (アプリが停止していた時間の割合)。", "Young領域サイズ (-Xmn/G1NewSizePercent) やオブジェクト割り当て速度のチューニング。"),
        ("▲ 注意", "GC 停止時間", "最長停止時間 >= 200ms", "最長 GC 停止が N ms に達しました (G1 目標値 200ms 超過)。", "-XX:MaxGCPauseMillis の調整、並行GCスレッド (ConcGCThreads) の見直し。"),
        ("▲ 注意", "スレッドプール / IO", "タスクスレッド使用率 >= 90%", "IO ワーカーのタスクスレッド使用率が最大 N% に達しました。", "HTTP処理スレッド枯渇の危機。task-max-threads 設定またはバックエンド呼出時間確認。"),
        ("▲ 注意", "CPU 資源不足", "実行可能スレッド > CPU数", "実行可能なスレッドの数 (95%値 N) が CPU 数 (M) を超えていました。", "CPU 不足によりスレッドが待たされています。タスクの vCPU 増強を推奨。"),
        ("▲ 注意", "CPU 資源不足", "PSI some avg10 >= 20%", "CPU を待った時間の割合 (PSI some avg10) が 95%値で N% ありました。", "カーネルレベルでの重大な CPU 競合。タスクの CPU 割り当て見直し。"),
        ("▲ 注意", "TCP 接続", "CLOSE_WAIT > 0", "CLOSE_WAIT が最大 N 本ありました。相手が切断したのに閉じていません。", "HTTPClient や RestTemplate のレスポンスストリーム close() 漏れを点検。"),
        ("● 参考", "収集の負荷", "常に評価", "収集処理 (jboss-cli/jstat/jcmd/perf-metrics) の CPU 使用率は平均 N%・最大 M% でした。", "本番との性能比較試験において収集オーバーヘッドが十分に軽微 (1〜2%) であることの確認。")
    ]

    for lv, cat, trig, msg, act in findings_rules:
        ws4.row_dimensions[r_idx].height = 28
        fill_cur = fill_row_even if r_idx % 2 == 0 else fill_row_odd
        
        c_lv = ws4.cell(row=r_idx, column=1, value=lv)
        c_lv.font = f_data_bold; c_lv.fill = fill_cur; c_lv.alignment = al_center; c_lv.border = bd_thin
        
        c_cat = ws4.cell(row=r_idx, column=2, value=cat)
        c_cat.font = f_data; c_cat.fill = fill_cur; c_cat.alignment = al_center; c_cat.border = bd_thin
        
        c_trig = ws4.cell(row=r_idx, column=3, value=trig)
        c_trig.font = f_cmd; c_trig.fill = fill_cur; c_trig.alignment = al_left; c_trig.border = bd_thin

        c_msg = ws4.cell(row=r_idx, column=4, value=msg)
        c_msg.font = f_data; c_msg.fill = fill_cur; c_msg.alignment = al_left; c_msg.border = bd_thin

        c_act = ws4.cell(row=r_idx, column=5, value=act)
        c_act.font = f_data; c_act.fill = fill_cur; c_act.alignment = al_left; c_act.border = bd_thin
        r_idx += 1

    r_idx += 1
    # 2. 全49グラフ
    ws4.cell(row=r_idx, column=1, value="2. レポート「グラフ」シートに出力される全 49 グラフの構成定義").font = f_section
    merge_and_style(ws4, r_idx, 1, r_idx, 7)
    r_idx += 1

    headers_ch = ["グラフ番号", "区分", "セクション / 分類", "グラフタイトル", "Y軸タイトル", "系列構成 (凡例項目)", "参照シート"]
    ws4.row_dimensions[r_idx].height = 24
    for ci, h in enumerate(headers_ch, start=1):
        cell = ws4.cell(row=r_idx, column=ci, value=h)
        cell.font = f_th; cell.fill = fill_th; cell.alignment = al_center; cell.border = bd_th
    r_idx += 1

    charts_def = [
        ("G01〜G03", "①", "① DB 接続プール", "<DS名>：接続数", "接続数", "使用中 (InUse), アクティブ (Active), 最大プール (Max)", "①DB接続プール"),
        ("G04〜G06", "①", "① DB 接続プール", "<DS名>：待ち・取得失敗（区間）", "件", "待ち (ΔWait), 取得失敗 (ΔBlockingFailure)", "①DB接続プール"),
        ("G07〜G09", "①", "① DB 接続プール", "<DS名>：平均待ちスレッド数・平均使用接続数", "スレッド / 接続", "平均待ちスレッド数 (TotalBlocking÷経過), 平均使用接続数", "①DB接続プール"),
        ("G10", "①", "① HTTP・セッション", "リクエスト 毎秒（リスナー別）", "件/秒", "リスナー別 RPS (default, https)", "①HTTP"),
        ("G11", "①", "① HTTP・セッション", "エラー（区間、リスナー別）", "件", "リスナー別 500系エラー件数", "①HTTP"),
        ("G12", "①", "① HTTP・セッション", "平均処理時間（ミリ秒、リスナー別）", "ミリ秒", "Undertow 内部平均処理時間", "①HTTP"),
        ("G13", "①", "① HTTP・セッション", "アクティブセッション数（デプロイメント別）", "個", "WAR アプリ別のアクティブ HTTP セッション数", "①セッション"),
        ("G14", "①", "① スレッドプール・IO", "IO ワーカー：タスクスレッド数", "スレッド", "使用中 (busy-task), 上限 (max-pool), コア (core-pool)", "①IOワーカー"),
        ("G15", "①", "① スレッドプール・IO", "IO ワーカー：キュー滞留数", "件", "タスクスレッド空き待ちキュー滞留数", "①IOワーカー"),
        ("G16", "①", "① スレッドプール・IO", "EE / EJB スレッドプール：実行中スレッド数", "スレッド", "プール別 実行中スレッド数 (active)", "①スレッドプール"),
        ("G17", "①", "① トランザクション", "トランザクション コミット 毎秒 (TPS)", "件/秒", "コミット毎秒 (TPS)", "①トランザクション"),
        ("G18", "①", "① トランザクション", "トランザクション 中断・タイムアウト（区間）", "件", "中断 (aborted), タイムアウト (timed-out)", "①トランザクション"),
        ("G19", "①", "① JVM（CLI）", "ヒープ使用量・確保量（MB、jboss-cli）", "MB", "ヒープ使用 (used), ヒープ確保 (committed), 上限 (max)", "①JVM（CLI）"),
        ("G20", "①", "① JVM（CLI）", "ライブスレッド数・OS スレッド数", "スレッド", "JVMライブスレッド, OSプロセススレッド (/proc/status)", "①JVM（CLI）"),
        ("G21", "②", "② ヒープ・GC（jstat）", "ヒープ各領域の使用量（MB、jstat）", "MB", "ヒープ合計, Eden使用, Survivor使用, Old使用", "②ヒープ・GC（jstat）"),
        ("G22", "②", "② ヒープ・GC（jstat）", "Old 領域使用率（%）", "%", "Old使用÷Old容量, Old使用÷最大ヒープ容量", "②ヒープ・GC（jstat）"),
        ("G23", "②", "② ヒープ・GC（jstat）", "GC 回数（区間、jstat）", "回", "Young GC 回数 (ΔYGC), Full GC 回数 (ΔFGC)", "②ヒープ・GC（jstat）"),
        ("G24", "②", "② ヒープ・GC（jstat）", "GC 停止時間（ミリ秒、区間、jstat）", "ミリ秒", "Young GC 時間 (ΔYGCT), Full GC 時間 (ΔFGCT)", "②ヒープ・GC（jstat）"),
        ("G25", "②", "② ヒープ・GC（jstat）", "GC 時間率（%、区間）", "%", "GC停止時間率 (ΔGCT÷経過秒)", "②ヒープ・GC（jstat）"),
        ("G26", "②", "② ヒープ・GC（jstat）", "セーフポイント停止時間（ミリ秒、区間）", "ミリ秒", "セーフポイント停止 (ΔSPT), 同期待ち (ΔSPS)", "②ヒープ・GC（jstat）"),
        ("G27", "②", "② ヒープ（jcmd）", "G1 リージョン数（jcmd）", "個", "Young リージョン数, Survivor リージョン数", "②ヒープ（jcmd）"),
        ("G28", "②", "② コードキャッシュ", "コードキャッシュ使用率（%）", "%", "セグメント別使用率 (Non-nmethods, Profiled, Non-profiled)", "②コードキャッシュ"),
        ("G29", "②", "② NMT", "ネイティブメモリ コミット量（MB、NMT）", "MB", "Total committed, Java Heap, Class, Thread, GC 等", "②NMT"),
        ("G30", "②", "② スレッド状態（jcmd）", "Java スレッド状態（全体、jcmd）", "スレッド", "処理中 (busy), 実行可能 (runnable), ブロック (blocked)", "②スレッド状態"),
        ("G31", "②", "② スレッド状態（jcmd）", "DB 接続待ちスレッド数（jcmd）", "スレッド", "DB接続待ちスレッド (db-pool-wait)", "②スレッド状態"),
        ("G32", "②", "② GC（JFR）", "GC ごとの最長停止時間（ミリ秒、JFR）", "ミリ秒", "GC最長停止時間 (longest_pause_ms)", "②GC（JFR）"),
        ("G33", "②", "② GC（JFR）", "GC 前後のヒープ使用量（MB、JFR）", "MB", "GC前ヒープ (heap_before), GC後ヒープ (heap_after)", "②GC（JFR）"),
        ("G34", "②", "② JVM統計（JFR）", "JVM CPU 使用率（%、JFR）", "%", "JVM User %, JVM System %, マシン全体 %", "②JVM統計（JFR）"),
        ("G35", "②", "② JVM統計（JFR）", "Java 例外発生 毎秒（JFR）", "件/秒", "例外発生レート (throwables/sec)", "②JVM統計（JFR）"),
        ("G36", "③", "③ top 概要", "ロードアベレージ（top）", "値", "1分 (load1), 5分 (load5), 15分 (load15)", "③top概要"),
        ("G37", "③", "③ top 概要", "CPU 内訳（%、全 CPU、top）", "%", "ユーザー (us), システム (sy), IO待ち (wa), steal (st)", "③top概要"),
        ("G38", "③", "③ top 概要", "プロセス別 CPU 使用率（%、1コア=100）", "%", "JBoss (jboss), 収集処理 (collect), その他 (other)", "③top概要"),
        ("G39", "③", "③ top スレッド", "JBoss CPU 上位スレッドの CPU 使用率", "%", "上位スレッド (TID・名前別) の CPU%", "③topスレッド"),
        ("G40", "③", "③ JBoss スレッド", "JBoss スレッド用途別 CPU（コア）", "コア", "HTTP処理 (http-task), GC (gc), JIT (jit), JFR (jfr)", "③JBossスレッド"),
        ("G41", "③", "③ JBoss スレッド", "JBoss 非自発的コンテキストスイッチ / 秒", "回/秒", "用途別 非自発的切替レート (all, http-task, gc)", "③JBossスレッド"),
        ("G42", "③", "③ FD・TCP", "ファイル記述子 使用数・上限", "個", "使用中 (fd_count), 上限 (fd_limit)", "③FD・TCP"),
        ("G43", "③", "③ FD・TCP", "TCP 接続数（JBoss）", "本", "ESTABLISHED (受信), ESTABLISHED (送信), CLOSE_WAIT", "③FD・TCP"),
        ("G44", "③", "③ FD・TCP", "TCP TIME_WAIT（タスク全体）", "本", "TIME_WAIT ソケット数", "③FD・TCP"),
        ("G45", "④", "④ CPU スケジューリング", "JBoss CPU 使用量 vs コンテナ取り分（コア）", "コア", "JBoss合計CPU, アプリCPU, GC CPU, 取り分(vCPU), JVM CPU数", "④CPUスケジューリング"),
        ("G46", "④", "④ CPU スケジューリング", "実行待ちスレッド数（procs_running）", "スレッド", "JBoss実行中スレッド, 実行可能スレッド全体, CPU数基準線", "④CPUスケジューリング"),
        ("G47", "④", "④ CPU スケジューリング", "CPU steal 率（%）", "%", "steal % (ほかのVMにCPUを奪われた時間割合)", "④CPUスケジューリング"),
        ("G48", "④", "④ CPU スケジューリング", "CPU の待ち（PSI some avg10 %）", "%", "PSI some avg10 % (CPU資源待機時間割合)", "④CPUスケジューリング"),
        ("G49", "④", "④ CPU スケジューリング", "cgroup CPU スロットリング（ミリ秒、区間）", "ミリ秒", "cfs_throttled_ms (クォータ制限時間)", "④CPUスケジューリング")
    ]

    for g_no, k, sec, title, y_title, series, sh_ref in charts_def:
        ws4.row_dimensions[r_idx].height = 22
        fill_cur = fill_row_even if r_idx % 2 == 0 else fill_row_odd
        
        c0 = ws4.cell(row=r_idx, column=1, value=g_no)
        c0.font = f_data_bold; c0.fill = fill_cur; c0.alignment = al_center; c0.border = bd_thin
        
        c1 = ws4.cell(row=r_idx, column=2, value=k)
        c1.font = f_badge; c1.fill = fill_badges.get(k, fill_th); c1.alignment = al_center; c1.border = bd_thin

        c2 = ws4.cell(row=r_idx, column=3, value=sec)
        c2.font = f_data; c2.fill = fill_cur; c2.alignment = al_left; c2.border = bd_thin

        c3 = ws4.cell(row=r_idx, column=4, value=title)
        c3.font = f_data_bold; c3.fill = fill_cur; c3.alignment = al_left; c3.border = bd_thin

        c4 = ws4.cell(row=r_idx, column=5, value=y_title)
        c4.font = f_data; c4.fill = fill_cur; c4.alignment = al_center; c4.border = bd_thin

        c5 = ws4.cell(row=r_idx, column=6, value=series)
        c5.font = f_data; c5.fill = fill_cur; c5.alignment = al_left; c5.border = bd_thin

        c6 = ws4.cell(row=r_idx, column=7, value=sh_ref)
        c6.font = f_data; c6.fill = fill_cur; c6.alignment = al_left; c6.border = bd_thin
        r_idx += 1

    auto_fit_columns(ws4, min_widths={'A': 12, 'B': 8, 'C': 20, 'D': 32, 'E': 16, 'F': 42, 'G': 20})

    # =============================================================
    # SHEET 5: ④環境設定・タスク定義 (Config & TaskDef)
    # =============================================================
    ws5 = wb.create_sheet(title="④環境設定・タスク定義")
    apply_title_banner(ws5, "環境変数パラメータ & ECS Fargate タスク定義仕様", 
                       "perf-metrics.sh の動作制御環境変数と性能試験用 ECS タスク定義 (JSON) の構成要件", 5)
    ws5.views.sheetView[0].showGridLines = True
    ws5.row_dimensions[3].height = 10

    # 1. 環境変数
    ws5.cell(row=4, column=1, value="1. perf-metrics.sh 収集制御環境変数一覧 (タスク定義 environment で上書き可能)").font = f_section
    merge_and_style(ws5, 4, 1, 4, 5)

    cols_env = ["環境変数名", "既定値", "設定可能値 / 候補", "説明・影響範囲", "推奨設定・備考"]
    r_idx = 5
    ws5.row_dimensions[r_idx].height = 24
    for ci, h in enumerate(cols_env, start=1):
        cell = ws5.cell(row=r_idx, column=ci, value=h)
        cell.font = f_th; cell.fill = fill_th; cell.alignment = al_center; cell.border = bd_th
    r_idx += 1

    env_vars = [
        ("PERF_METRICS_ENABLED", "false", "true | false", "性能試験メトリクス収集を有効化するか。エントリポイントが判定。", "Dockerfile ビルド引数で指定。本番は false で完全除去。"),
        ("PERF_INTERVAL", "10", "正の整数 (秒)", "各コレクター (cli, jstat, top, ps, proc, cpu) の既定の収集間隔。", "基本 10 秒を推奨。高精度測定時は 5 秒。"),
        ("PERF_CLI_INTERVAL", "${PERF_INTERVAL}", "正の整数 (秒)", "① jboss-cli による DB/HTTP/スレッドプールの収集間隔。", "既定 (10秒) のままで十分。"),
        ("PERF_JSTAT_INTERVAL", "${PERF_INTERVAL}", "正の整数 (秒)", "② jstat (perfgc) によるヒープ・GC・セーフポイントの収集間隔。", "1 プロセス常駐のため 5〜10 秒でも極めて軽量。"),
        ("PERF_JCMD_INTERVAL", "60", "正の整数 (秒)", "② jcmd による GC.heap_info, Compiler.codecache, Thread.print の間隔。", "1回ごとに JVM を起動するため 60 秒以上を強く推奨。"),
        ("PERF_JFR", "true", "true | false", "Java Flight Recorder (JFR) を JBoss 起動時から常時記録するか。", "true を強く推奨。試験後の障害解析に絶大な威力。"),
        ("PERF_JFR_SETTINGS", "heap", "heap | default | profile | パス", "JFR のイベント設定プロファイル。heap はメモリ・GC・統計に特化。", "CPU プロファイルも詳細に取る場合は profile を指定。"),
        ("PERF_JFR_MAXAGE", "24h", "時間文字列 (例: 1h, 24h)", "JFR のディスク保持期間上限。", "試験シナリオの長さに応じて設定。"),
        ("PERF_JFR_MAXSIZE", "1g", "サイズ文字列 (例: 500m, 1g)", "JFR ファイルの最大サイズ上限。超過時は古いチャンクから自動削除。", "EFS の容量を考慮して設定。"),
        ("PERF_COLLECTORS", "cli jstat jcmd jfr cpu", "空白区切り (none で全停止)", "稼働させる収集コレクターのリスト。", "特定のコレクターだけを動かしたい場合に指定。"),
        ("PERF_ENABLE_STATISTICS", "true", "true | false", "起動引数に各サブシステムの statistics-enabled=true を足すか。", "true で DB・Undertow・TX の統計フラグを自動付与。"),
        ("PERF_NMT", "off", "off | summary | detail", "Native Memory Tracking を有効化するか (jcmd VM.native_memory)。", "JVMが数%重くなるため、オフヒープリーク調査時のみ summary を推奨。"),
        ("PERF_PROCPS", "auto", "auto | true | false", "Dockerfile ビルド時に microdnf で procps-ng (top/ps) を導入するか。", "auto で導入を試み、失敗しても /proc 代替で動作。"),
        ("PERF_JVM_CONTAINERS", "(自動判定)", "正の整数", "④ の判定でタスク内の JVM 稼働コンテナ数。サイドカー除外後の数。", "自動判定で意図通りにならない場合に明示。")
    ]

    for ev, dflt, vals, desc, note in env_vars:
        ws5.row_dimensions[r_idx].height = 24
        fill_cur = fill_row_even if r_idx % 2 == 0 else fill_row_odd
        
        c0 = ws5.cell(row=r_idx, column=1, value=ev)
        c0.font = f_cmd; c0.fill = fill_cur; c0.alignment = al_left; c0.border = bd_thin
        
        c1 = ws5.cell(row=r_idx, column=2, value=dflt)
        c1.font = f_data; c1.fill = fill_cur; c1.alignment = al_center; c1.border = bd_thin

        c2 = ws5.cell(row=r_idx, column=3, value=vals)
        c2.font = f_data; c2.fill = fill_cur; c2.alignment = al_left; c2.border = bd_thin

        c3 = ws5.cell(row=r_idx, column=4, value=desc)
        c3.font = f_data; c3.fill = fill_cur; c3.alignment = al_left; c3.border = bd_thin

        c4 = ws5.cell(row=r_idx, column=5, value=note)
        c4.font = f_data; c4.fill = fill_cur; c4.alignment = al_left; c4.border = bd_thin
        r_idx += 1

    r_idx += 1
    # 2. タスク定義
    ws5.cell(row=r_idx, column=1, value="2. 性能試験用 ECS Fargate タスク定義 (taskdef-perf.jsonc) の必須設定要件").font = f_section
    merge_and_style(ws5, r_idx, 1, r_idx, 5)
    r_idx += 1

    cols_task = ["設定項目", "本番環境との差異", "推奨設定内容", "設定理由・技術的背景", "未設定時のリスク・影響"]
    ws5.row_dimensions[r_idx].height = 24
    for ci, h in enumerate(cols_task, start=1):
        cell = ws5.cell(row=r_idx, column=ci, value=h)
        cell.font = f_th; cell.fill = fill_th; cell.alignment = al_center; cell.border = bd_th
    r_idx += 1

    taskdef_reqs = [
        ("/tmp ボリュームマウント\n(エフェメラル)", "本番: readonly または共用\n性能試験: 各コンテナ個別ボリューム",
         '{\n  "name": "front-tmp",\n  "dockerVolumeConfiguration": {}\n}\nコンテナ側: "mountPoints": [{"sourceVolume": "front-tmp", "containerPath": "/tmp"}]',
         "jstat / jcmd が JVM にアタッチするための hsperfdata_<user>/<pid> およびソケットファイルを配置するため。\n※フロントとバックで必ず別ボリュームにすること (PID 重複による誤認防止)。",
         "jstat/jcmd が一切接続できず、ヒープ詳細・スレッド状態が採取不可になる。"),
        ("メモリ割り当て余裕\n(Memory Reservation)", "本番: JBoss 専用サイズ\n性能試験: +250MB〜300MB 確保",
         'タスクメモリおよびコンテナの memoryReservation に 250MB 程度の余裕を確保。\n(例: 本番 4096MB → 試験用 4352MB)',
         "収集用の jboss-cli (常駐64MB), jstat (常駐32MB), jcmd (一時起動), JFR CSV変換等の RSS 合計が 200〜250MB 程度増加するため。",
         "Fargate の OOM Killer により JBoss または収集プロセスが強制停止される。"),
        ("コンテナの `cpu` 指定\n(CPU ユニット数)", "本番: 未指定 (タスク按分)\n性能試験: 各コンテナに明示指定",
         '{\n  "name": "intra-web-front",\n  "cpu": 1024  // 1 vCPU = 1024 ユニット\n}',
         "④ の判定ロジックが「このコンテナに割り当てられた取り分 (vCPU)」を誤認なく一意に確定できるようにするため。",
         "タスクの vCPU ÷ JVM コンテナ数 の推測計算にフォールバックされる。"),
        ("ECS Exec 有効化\n& タスクロール権限", "本番: 原則無効\n性能試験: 有効化推奨",
         'タスクロールに "ssmmessages:*" 権限を付与。\naws ecs run-task / update-service で --enable-execute-command を付与。',
         "試験中にコンテナへ入って `perf-metrics.sh status`, `threaddump`, `cpu-check`, `pack` などの対話操作を行うため。",
         "試験中の動的なスレッドダンプ取得や JFR ダンプ、状態確認ができなくなる。")
    ]

    for req, diff, val, reason, risk in taskdef_reqs:
        ws5.row_dimensions[r_idx].height = 64
        fill_cur = fill_row_even if r_idx % 2 == 0 else fill_row_odd
        
        c0 = ws5.cell(row=r_idx, column=1, value=req)
        c0.font = f_data_bold; c0.fill = fill_cur; c0.alignment = al_center; c0.border = bd_thin
        
        c1 = ws5.cell(row=r_idx, column=2, value=diff)
        c1.font = f_data; c1.fill = fill_cur; c1.alignment = al_left; c1.border = bd_thin

        c2 = ws5.cell(row=r_idx, column=3, value=val)
        c2.font = f_cmd; c2.fill = fill_cur; c2.alignment = al_left; c2.border = bd_thin

        c3 = ws5.cell(row=r_idx, column=4, value=reason)
        c3.font = f_data; c3.fill = fill_cur; c3.alignment = al_left; c3.border = bd_thin

        c4 = ws5.cell(row=r_idx, column=5, value=risk)
        c4.font = f_data; c4.fill = fill_cur; c4.alignment = al_left; c4.border = bd_thin
        r_idx += 1

    auto_fit_columns(ws5, min_widths={'A': 22, 'B': 20, 'C': 34, 'D': 38, 'E': 32})

    # デフォルトシート削除
    wb.remove(default_sheet)

    # 確実に全セルを Meiryo UI に統一
    enforce_meiryo_ui(wb)

    # 保存
    wb.save(output_path)
    print(f"Workbook successfully saved to: {output_path}")

if __name__ == "__main__":
    out_file = "JMeterメトリクス収集_レポート出力項目と実装コマンド一覧.xlsx"
    build_excel(out_file)
